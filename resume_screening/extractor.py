"""Turn an uploaded resume (PDF / DOCX / DOC / TXT) into plain text."""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from .text_utils import clean_whitespace

MIN_TEXT_CHARS = 40


class ExtractionError(ValueError):
    """Raised with a user-presentable message when a file cannot be read."""


def extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def extract_text(filename: str, data: bytes) -> str:
    ext = extension(filename)
    if not data:
        raise ExtractionError("The file is empty.")
    if ext == "pdf":
        text = _from_pdf(data)
    elif ext == "docx":
        text = _from_docx(data)
    elif ext == "doc":
        text = _from_doc(data)
    elif ext == "txt":
        text = _from_txt(data)
    else:
        raise ExtractionError(f"Unsupported file type '.{ext}'. Upload PDF, DOC, DOCX or TXT.")
    text = clean_whitespace(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "No readable text found. Scanned/image-only files are not supported - "
            "export the resume as a text-based PDF or DOCX."
        )
    return text


def _from_txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("latin-1", errors="ignore")


def _from_pdf(data: bytes) -> str:
    if not data.lstrip()[:5] == b"%PDF-":
        raise ExtractionError("This file is not a valid PDF.")
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("The PDF is password-protected.")
            except ExtractionError:
                raise
            except Exception:
                raise ExtractionError("The PDF is password-protected.")
        pages = []
        for page in reader.pages[:15]:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    except ExtractionError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        raise ExtractionError(f"Could not read the PDF ({exc.__class__.__name__}).") from exc


def _from_docx(data: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractionError("This file is not a valid DOCX document.")
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises assorted errors on bad packages
        raise ExtractionError("Could not read the DOCX document.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                parts.append(cell.text)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
    return "\n".join(parts)


OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _from_doc(data: bytes) -> str:
    """Legacy Word 97-2003 files: antiword when installed, else a best-effort
    scan of the WordDocument stream."""
    if not data.startswith(OLE_MAGIC):
        if zipfile.is_zipfile(io.BytesIO(data)):  # a .docx saved with a .doc extension
            return _from_docx(data)
        raise ExtractionError("This file is not a valid DOC document.")
    antiword = shutil.which("antiword")
    if antiword:
        with tempfile.NamedTemporaryFile(suffix=".doc") as tmp:
            tmp.write(data)
            tmp.flush()
            try:
                out = subprocess.run(
                    [antiword, "-m", "UTF-8.txt", tmp.name],
                    capture_output=True, timeout=30, check=False,
                )
                text = out.stdout.decode("utf-8", errors="ignore")
                if len(text.strip()) >= MIN_TEXT_CHARS:
                    return text
            except (subprocess.TimeoutExpired, OSError):
                pass
    return _doc_fallback(data)


def _doc_fallback(data: bytes) -> str:
    try:
        import olefile
    except ImportError:  # pragma: no cover
        raise ExtractionError("DOC support requires antiword or olefile.")
    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
        stream = ole.openstream("WordDocument").read()
    except Exception as exc:
        raise ExtractionError("Could not read the DOC document.") from exc
    # Text is stored either as 8-bit cp1252 or UTF-16LE; scan for both.
    ascii_runs = re.findall(rb"[\x20-\x7e\r\n\t\x07\x0b\x0c\x91-\x97]{4,}", stream)
    ascii_text = b"\n".join(ascii_runs).decode("cp1252", errors="ignore")
    utf16_runs = re.findall(rb"(?:[\x20-\x7e\r\n\t]\x00){4,}", stream)
    utf16_text = "\n".join(r.decode("utf-16le", errors="ignore") for r in utf16_runs)
    text = utf16_text if len(utf16_text) > len(ascii_text) else ascii_text
    return text.replace("\r", "\n").replace("\x07", " ").replace("\x0b", "\n").replace("\x0c", "\n")
