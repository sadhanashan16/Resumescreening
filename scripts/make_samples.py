"""Regenerate the demo resumes in samples/ (PDF, DOCX, DOC, TXT).

    pip install -r requirements-dev.txt
    python scripts/make_samples.py

All people are fictional (generated names, example.com e-mail addresses).
The .doc sample is produced from a .docx with LibreOffice when it is installed.
"""
from __future__ import annotations

import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from resume_screening.dataset import generate_resume  # noqa: E402

OUT = ROOT / "samples"
# (role, seed, format)
SAMPLES = [
    ("data-scientist", 11, "pdf"), ("machine-learning-engineer", 23, "docx"), ("data-analyst", 5, "pdf"),
    ("data-engineer", 8, "txt"), ("backend-developer", 31, "docx"), ("frontend-developer", 14, "pdf"),
    ("full-stack-developer", 2, "docx"), ("devops-engineer", 5, "pdf"), ("mobile-app-developer", 9, "txt"),
    ("qa-engineer", 17, "doc"), ("cybersecurity-analyst", 4, "pdf"), ("business-analyst", 6, "docx"),
    ("data-scientist", 40, "txt"), ("frontend-developer", 3, "docx"),
]


def write_txt(r, path):
    path.write_text(r.to_text(), encoding="utf-8")


def write_docx(r, path):
    from docx import Document
    from docx.shared import Pt

    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading(r.name, level=0)
    d.add_paragraph(f"{r.email} | {r.phone}")
    for heading, lines in r.sections:
        d.add_heading(heading, level=1)
        for line in lines:
            if line.startswith("- "):
                d.add_paragraph(line[2:], style="List Bullet")
            else:
                d.add_paragraph(line)
    d.save(path)


def write_pdf(r, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], fontSize=22, alignment=0, spaceAfter=2)
    head = ParagraphStyle("h", parent=ss["Heading2"], fontSize=12, textColor=colors.HexColor("#3730a3"), spaceBefore=10, spaceAfter=3)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"{r.name} - Resume", author=r.name)
    story = [Paragraph(escape(r.name), title), Paragraph(escape(f"{r.email} | {r.phone}"), body)]
    for heading, lines in r.sections:
        story.append(Paragraph(escape(heading), head))
        bullets = []
        for line in lines:
            if line.startswith("- "):
                bullets.append(ListItem(Paragraph(escape(line[2:]), body), leftIndent=12))
            else:
                if bullets:
                    story.append(ListFlowable(bullets, bulletType="bullet", start="•", leftIndent=12)); bullets = []
                story.append(Paragraph(escape(line), body))
        if bullets:
            story.append(ListFlowable(bullets, bulletType="bullet", start="•", leftIndent=12))
        story.append(Spacer(1, 2))
    doc.build(story)


def write_doc(r, path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print(f"  ! LibreOffice not found - skipping {path.name}")
        return False
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / (path.stem + ".docx")
        write_docx(r, src)
        subprocess.run([soffice, "--headless", "--convert-to", "doc", "--outdir", tmp, str(src)],
                       check=True, capture_output=True, timeout=120)
        shutil.copy(Path(tmp) / (path.stem + ".doc"), path)
    return True


def main():
    OUT.mkdir(exist_ok=True)
    for old in OUT.iterdir():
        if old.is_file():
            old.unlink()
    writers = {"txt": write_txt, "docx": write_docx, "pdf": write_pdf, "doc": write_doc}
    for role, seed, fmt in SAMPLES:
        r = generate_resume(role, random.Random(seed), hybrid=False)
        path = OUT / f"{r.name.replace(' ', '_')}_{role.replace('-', '_')}.{fmt}"
        writers[fmt](r, path)
        print("wrote", path.name)


if __name__ == "__main__":
    main()
